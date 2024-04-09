import os
import datetime
from fpdf import FPDF
from docx import Document
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart

samples = []

class PDF(FPDF):
    def header(self):
        self.set_font('Arial', 'B', 12)
        self.cell(0, 10, 'Мой заголовок', 0, 1, 'C')

    def footer(self):
        self.set_y(-15)
        self.set_font('Arial', 'I', 8)
        self.cell(0, 10, f'Страница {self.page_no()}', 0, 0, 'C')

def create_letter(sample):
    return f"{sample['text']}\nДата создания: {sample['date_created']}"

def save_letter_as_pdf(letter, filename):
    pdf = PDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.multi_cell(0, 10, txt=letter, align='L')
    pdf.output(filename)

def save_letter_as_docx(letter, filename):
    doc = Document()
    doc.add_paragraph(letter)
    doc.save(filename)

def save_letter_as_txt(letter, filename):
    with open(filename, "w", encoding="utf-8") as file:
        file.write(letter)

def print_saved_files():
    saved_files = [f for f in os.listdir() if f.endswith((".pdf", ".docx", ".txt"))]
    if not saved_files:
        print("Нет сохраненных документов.")
    else:
        print("Список сохраненных документов:")
        for index, filename in enumerate(saved_files, start=1):
            print(f"{index}. {filename}")

        file_index = input("Выберите номер файла для печати (или 0 для возврата): ")
        if file_index == "0":
            return  # Возврат к главному меню
        try:
            file_index = int(file_index)
            if 0 < file_index <= len(saved_files):
                selected_file = saved_files[file_index - 1]
                print_file(selected_file)
            else:
                print("Некорректный номер файла.")
        except ValueError:
            print("Некорректный ввод. Введите номер файла для печати или 0 для возврата.")

def print_file(filename):
    try:
        with open(filename, "r", encoding="utf-8") as file:
            print(file.read())
    except FileNotFoundError:
        print("Файл не найден.")


def send_email(subject, body, to_email):
    from_email = "zatura55@mail.ru"  # Замените на свой email
    password = "fzUP5CYJQyLcXmdynwfn"  # Замените на пароль от вашего email

    msg = MIMEMultipart()
    msg["From"] = from_email
    msg["To"] = to_email
    msg["Subject"] = subject

    msg.attach(MIMEText(body, "plain"))

    try:
        server = smtplib.SMTP("smtp.mail.ru", 587)
        server.starttls()
        server.login(from_email, password)
        text = msg.as_string()
        server.sendmail(from_email, to_email, text)
        server.quit()
        print(f"Письмо с темой '{subject}' отправлено на адрес {to_email}.")
    except Exception as e:
        print(f"Ошибка при отправке письма: {e}")

def caesar_cipher(text, shift):
    encrypted_text = ""
    for char in text:
        if char.isalpha():
            is_upper = char.isupper()
            char = char.lower()
            shifted_char = chr(((ord(char) - ord('а') + shift) % 32) + ord('а'))
            if is_upper:
                shifted_char = shifted_char.upper()
            encrypted_text += shifted_char
        else:
            encrypted_text += char
    return encrypted_text

def caesar_decipher(text, shift):
    decrypted_text = ""
    for char in text:
        if char.isalpha():
            is_upper = char.isupper()
            char = char.lower()
            shifted_char = chr(((ord(char) - ord('а') - shift) % 32) + ord('а'))
            if is_upper:
                shifted_char = shifted_char.upper()
            decrypted_text += shifted_char
        else:
            decrypted_text += char
    return decrypted_text

while True:
    print("1. Создать новый образец")
    print("2. Просмотреть существующие образцы")
    print("3. Редактировать существующий образец")
    print("4. Сохранить текущее письмо на диск")
    print("5. Сохранить существующий образец на диск")
    print("6. Отправить письмо по электронной почте")
    print("7. Печать сохраненных документов")
    print("8. Печать выбранного документа")
    print("9. Зашифровать текст шифром Цезаря")
    print("10. Расшифровать текст шифром Цезаря")
    print("0. Выход")

    choice = input("Выберите действие (0/1/2/3/4/5/6/7/8/9/10): ")

    if choice == "1":
        # Создание нового образца
        text = input("Введите текст образца ответа:\n")
        sample = {
            "text": text,
            "date_created": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }
        samples.append(sample)
        print("Образец успешно создан.")

    elif choice == "2":
        # Просмотр существующих образцов
        if not samples:
            print("Список образцов ответов пуст.")
        else:
            print("Список существующих образцов:")
            for index, sample in enumerate(samples, start=1):
                print(f"{index}. {sample['text']}")
            print("Выберите номер образца для просмотра (или 0 для возврата): ")
            sample_index = int(input())
            if sample_index > 0 and sample_index <= len(samples):
                print(create_letter(samples[sample_index - 1]))
            elif sample_index != 0:
                print("Некорректный номер образца.")

    elif choice == "3":
        # Редактирование существующего образца
        if not samples:
            print("Список образцов ответов пуст.")
        else:
            print("Список существующих образцов:")
            for index, sample in enumerate(samples, start=1):
                print(f"{index}. {sample['text']}")
            print("Выберите номер образца для редактирования (или 0 для возврата): ")
            sample_index = int(input())
            if sample_index > 0 and sample_index <= len(samples):
                new_text = input("Введите новый текст образца ответа:\n")
                samples[sample_index - 1]["text"] = new_text
                samples[sample_index - 1]["date_created"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                print("Образец успешно отредактирован.")
            elif sample_index != 0:
                print("Некорректный номер образца.")

    elif choice == "4":
        # Сохранение текущего письма на диск
        if not samples:
            print("Список образцов ответов пуст.")
        else:
            print("Список существующих образцов:")
            for index, sample in enumerate(samples, start=1):
                print(f"{index}. {sample['text']}")
            print("Выберите номер образца для сохранения на диск (или 0 для возврата): ")
            sample_index = int(input())
            if sample_index > 0 and sample_index <= len(samples):
                selected_sample = samples[sample_index - 1]
                letter = create_letter(selected_sample)
                print("\nВыберите формат сохранения:")
                print("1. PDF")
                print("2. DOCX")
                print("3. TXT")
                format_choice = input("Выберите формат (1/2/3): ")
                filename = input("Введите имя файла для сохранения (без расширения): ")
                if format_choice == "1":
                    save_letter_as_pdf(letter, f"{filename}.pdf")
                    print(f"Образец успешно сохранен в файл {filename}.pdf")
                elif format_choice == "2":
                    save_letter_as_docx(letter, f"{filename}.docx")
                    print(f"Образец успешно сохранен в файл {filename}.docx")
                elif format_choice == "3":
                    save_letter_as_txt(letter, f"{filename}.txt")
                    print(f"Образец успешно сохранен в файл {filename}.txt")
                else:
                    print("Некорректный выбор формата.")
            elif sample_index != 0:
                print("Некорректный номер образца.")

    elif choice == "5":
        # Сохранение существующего образца на диск
        if not samples:
            print("Список образцов ответов пуст.")
        else:
            print("Список существующих образцов:")
            for index, sample in enumerate(samples, start=1):
                print(f"{index}. {sample['text']}")
            print("Выберите номер образца для сохранения на диск (или 0 для возврата): ")
            sample_index = int(input())
            if sample_index > 0 and sample_index <= len(samples):
                selected_sample = samples[sample_index - 1]
                letter = create_letter(selected_sample)
                print("\nВыберите формат сохранения:")
                print("1. PDF")
                print("2. DOCX")
                print("3. TXT")
                format_choice = input("Выберите формат (1/2/3): ")
                filename = input("Введите имя файла для сохранения (без расширения): ")
                if format_choice == "1":
                    save_letter_as_pdf(letter, f"{filename}.pdf")
                    print(f"Образец успешно сохранен в файл {filename}.pdf")
                elif format_choice == "2":
                    save_letter_as_docx(letter, f"{filename}.docx")
                    print(f"Образец успешно сохранен в файл {filename}.docx")
                elif format_choice == "3":
                    save_letter_as_txt(letter, f"{filename}.txt")
                    print(f"Образец успешно сохранен в файл {filename}.txt")
                else:
                    print("Некорректный выбор формата.")
            elif sample_index != 0:
                print("Некорректный номер образца.")

    elif choice == "6":
        # Отправка письма по электронной почте
        if not samples:
            print("Список образцов ответов пуст.")
        else:
            print("Список существующих образцов:")
            for index, sample in enumerate(samples, start=1):
                print(f"{index}. {sample['text']}")
            print("Выберите номер образца для отправки по электронной почте (или 0 для возврата): ")
            sample_index = int(input())
            if sample_index > 0 and sample_index <= len(samples):
                selected_sample = samples[sample_index - 1]
                letter = create_letter(selected_sample)
                to_email = input("Введите адрес электронной почты получателя: ")
                subject = input("Введите тему письма: ")
                send_email(subject, letter, to_email)
            elif sample_index != 0:
                print("Некорректный номер образца.")

    elif choice == "7":
        # Печать сохраненных документов
        saved_files = [f for f in os.listdir() if f.endswith((".pdf", ".docx", ".txt"))]
        if not saved_files:
            print("Нет сохраненных документов.")
        else:
            print("Список сохраненных PDF-документов:")
            pdf_files = [f for f in saved_files if f.endswith(".pdf")]
            for index, pdf_file in enumerate(pdf_files, start=1):
                print(f"{index}. {pdf_file}")
            print("Список сохраненных DOCX-документов:")
            docx_files = [f for f in saved_files if f.endswith(".docx")]
            for index, docx_file in enumerate(docx_files, start=len(pdf_files) + 1):
                print(f"{index}. {docx_file}")
            print("Список сохраненных TXT-документов:")
            txt_files = [f for f in saved_files if f.endswith(".txt")]
            for index, txt_file in enumerate(txt_files, start=len(pdf_files) + len(docx_files) + 1):
                print(f"{index}. {txt_file}")
            print("Выберите номер файла для печати (или 0 для возврата): ")
            file_index = int(input())
            if file_index > 0 and file_index <= len(saved_files):
                print_file(saved_files[file_index - 1])
            elif file_index != 0:
                print("Некорректный номер файла.")

    elif choice == "8":
        # Печать выбранного документа
        saved_files = [f for f in os.listdir() if f.endswith((".pdf", ".docx", ".txt"))]
        if not saved_files:
            print("Нет сохраненных документов.")
        else:
            print("Список сохраненных PDF-документов:")
            pdf_files = [f for f in saved_files if f.endswith(".pdf")]
            for index, pdf_file in enumerate(pdf_files, start=1):
                print(f"{index}. {pdf_file}")
            print("Список сохраненных DOCX-документов:")
            docx_files = [f for f in saved_files if f.endswith(".docx")]
            for index, docx_file in enumerate(docx_files, start=len(pdf_files) + 1):
                print(f"{index}. {docx_file}")
            print("Список сохраненных TXT-документов:")
            txt_files = [f for f in saved_files if f.endswith(".txt")]
            for index, txt_file in enumerate(txt_files, start=len(pdf_files) + len(docx_files) + 1):
                print(f"{index}. {txt_file}")
            print("Выберите номер файла для печати (или 0 для возврата): ")
            file_index = int(input())
            if file_index > 0 and file_index <= len(saved_files):
                print_file(saved_files[file_index - 1])
            elif file_index != 0:
                print("Некорректный номер файла.")

    elif choice == "9":
        # Зашифровать текст шифром Цезаря
        text_to_cipher = input("Введите текст для шифрования шифром Цезаря (на русском): ")
        shift = int(input("Введите сдвиг: "))
        encrypted_text = caesar_cipher(text_to_cipher, shift)
        print(f"Зашифрованный текст: {encrypted_text}")

    elif choice == "10":
        # Расшифровать текст шифром Цезаря
        text_to_decipher = input("Введите текст для расшифровки шифром Цезаря (на русском): ")
        shift = int(input("Введите сдвиг: "))
        decrypted_text = caesar_decipher(text_to_decipher, shift)
        print(f"Расшифрованный текст: {decrypted_text}")

    elif choice == "0":
        # Выход из программы
        print("Программа завершена.")
        break

    else:
        print("Некорректный выбор. Пожалуйста, выберите действие (0/1/2/3/4/5/6/7/8/9/10).")
