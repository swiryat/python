import os
import datetime
from fpdf import FPDF
from docx import Document

# Список образцов ответов
samples = []

# Класс для создания PDF
class PDF(FPDF):
    def header(self):
        self.set_font('Arial', 'B', 12)
        self.cell(0, 10, 'Мой заголовок', 0, 1, 'C')

    def footer(self):
        self.set_y(-15)
        self.set_font('Arial', 'I', 8)
        self.cell(0, 10, f'Страница {self.page_no()}', 0, 0, 'C')

def create_letter(sample):
    # Генерация письма на основе образца
    return f"{sample['text']}\nДата создания: {sample['date_created']}"

def save_letter_as_pdf(letter, filename):
    # Сохранение письма в PDF
    pdf = PDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.multi_cell(0, 10, txt=letter, align='L')
    pdf.output(filename)

def save_letter_as_docx(letter, filename):
    # Сохранение письма в DOCX
    doc = Document()
    doc.add_paragraph(letter)
    doc.save(filename)

def save_letter_as_txt(letter, filename):
    # Сохранение письма в TXT
    with open(filename, "w", encoding="utf-8") as file:
        file.write(letter)

def print_file(filename):
    # Печать содержимого файла
    try:
        with open(filename, "r", encoding="utf-8") as file:
            print(file.read())
    except FileNotFoundError:
        print("Файл не найден.")

def send_email(subject, body, to_email):
    # Здесь можно добавить код для отправки письма по электронной почте
    print(f"Письмо с темой '{subject}' отправлено на адрес {to_email}:\n{body}")

while True:
    print("1. Создать новый образец")
    print("2. Просмотреть существующие образцы")
    print("3. Редактировать существующий образец")
    print("4. Сохранить текущее письмо на диск")
    print("5. Сохранить существующий образец на диск")
    print("6. Отправить письмо по электронной почте")
    print("7. Печать сохраненных документов")
    print("8. Печать выбранного документа")
    print("9. Выход")
    
    choice = input("Выберите действие (1/2/3/4/5/6/7/8/9): ")
    
    if choice == "1":
        # Создание нового образца
        text = input("Введите текст образца ответа:\n")
        sample = {
            "text": text,
            "date_created": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }
        samples.append(sample)
        print("Образец успешно создан.")
    elif choice == "2":
        # Просмотр существующих образцов
        if not samples:
            print("Список образцов ответов пуст.")
        else:
            print("Список существующих образцов:")
            for index, sample in enumerate(samples, start=1):
                print(f"{index}. {sample['text']}")
            print("Выберите номер образца для просмотра (или 0 для возврата): ")
            sample_index = int(input())
            if sample_index > 0 and sample_index <= len(samples):
                print(create_letter(samples[sample_index - 1]))
            elif sample_index != 0:
                print("Некорректный номер образца.")
    elif choice == "3":
        # Редактирование существующего образца
        if not samples:
            print("Список образцов ответов пуст.")
        else:
            print("Список существующих образцов:")
            for index, sample in enumerate(samples, start=1):
                print(f"{index}. {sample['text']}")
            print("Выберите номер образца для редактирования (или 0 для возврата): ")
            sample_index = int(input())
            if sample_index > 0 and sample_index <= len(samples):
                new_text = input("Введите новый текст образца ответа:\n")
                samples[sample_index - 1]["text"] = new_text
                samples[sample_index - 1]["date_created"] = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                print("Образец успешно отредактирован.")
            elif sample_index != 0:
                print("Некорректный номер образца.")
    elif choice == "4":
        # Сохранение текущего письма на диск
        if not samples:
            print("Список образцов ответов пуст.")
        else:
            print("Список существующих образцов:")
            for index, sample in enumerate(samples, start=1):
                print(f"{index}. {sample['text']}")
            print("Выберите номер образца для сохранения на диск (или 0 для возврата): ")
            sample_index = int(input())
            if sample_index > 0 and sample_index <= len(samples):
                selected_sample = samples[sample_index - 1]
                letter = create_letter(selected_sample)
                print("\nВыберите формат сохранения:")
                print("1. PDF")
                print("2. DOCX")
                print("3. TXT")
                format_choice = input("Выберите формат (1/2/3): ")
                filename = input("Введите имя файла для сохранения (без расширения): ")
                if format_choice == "1":
                    save_letter_as_pdf(letter, f"{filename}.pdf")
                    print(f"Образец успешно сохранен в файл {filename}.pdf")
                elif format_choice == "2":
                    save_letter_as_docx(letter, f"{filename}.docx")
                    print(f"Образец успешно сохранен в файл {filename}.docx")
                elif format_choice == "3":
                    save_letter_as_txt(letter, f"{filename}.txt")
                    print(f"Образец успешно сохранен в файл {
